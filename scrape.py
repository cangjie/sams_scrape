#!/usr/local/bin/python3
"""山姆 App 商品列表抓取：iPhone 镜像窗口截图 → 本机 OCR → 逐个子分类翻到底。

只做界面操作：点顶部大类、点左侧子分类、滚动商品列表。不点商品、不点加购。

用法（/usr/local/bin/python3，装有 openpyxl）：
  scrape.py run [--top 大类名] [--max-pages N] [--wait-minutes 30]
        抓当前（或指定）大类；有断点就从断点所在的子分类重新扫一遍接着抓。
  scrape.py status            看进度
  scrape.py export            由 data/items.jsonl 生成 Excel
  scrape.py replay <截图目录>  用已有截图离线重跑解析（调试解析规则用）

暂停方式（任选其一，进度每页落盘）：
  - 直接动鼠标：脚本发现光标被挪开就停；
  - touch data/STOP；
  - 来电话、弹窗等导致页面不是「分类」页：脚本原地等待、不碰鼠标。页面恢复且镜像窗口在最前时
    自动继续；超过 --wait-minutes 仍未恢复则退出。
再次执行 run 即从断点继续，已抓到的商品自动去重。

数据都在 data/：state.json（断点）、items.jsonl（逐条追加）、shots/<大类>/（截图）、山姆商品_暂存.xlsx。
"""
import argparse
import difflib
import json
import os
import platform
import random
import re
import subprocess
import sys
import time
from datetime import datetime
from pathlib import Path

ROOT = Path(__file__).resolve().parent
BIN = ROOT / "bin"
# 识别程序：Mac mini（arm64）用 bin/analyze；Intel 笔记本用 bin/analyze-x86_64（用 src/analyze.swift 在本机编译）
ANALYZE = BIN / ("analyze-x86_64" if platform.machine() == "x86_64" else "analyze")
DATA = ROOT / "data"
STATE = DATA / "state.json"
ITEMS = DATA / "items.jsonl"
XLSX = DATA / "山姆商品_暂存.xlsx"
STOP = DATA / "STOP"

BASE_W, BASE_H = 383.0, 834.0        # 布局常量按这个窗口尺寸定义
SIDEBAR_X_MAX = 95                   # 左侧子分类栏的右边界
IMG_X_MIN = 100                      # 商品图左边界
LIST_X_MIN, LIST_X_MAX = 185, 320    # 商品文字区（左边是商品图，右边是加购按钮/悬浮按钮）
SKIP_SUBS = {"为您推荐", "新品上市"}   # 跨分类的推荐列表，不单独抓
TAB_WORDS = ("首页", "分类", "发现", "购物车", "我的")
SCROLL_DY = -255                     # 每页滚动像素；要小于「可见高度 - 一张商品卡」，保证每张卡至少完整出现一次

PRICE_RE = re.compile(r"^[¥￥羊半vVyY]?\s*(\d+(?:\.\d+)?)")
UNITS = r"(?:kg|千克|公斤|克|g|毫升|ml|升|l|斤|磅|片|只|个(?!月)|枚|粒|包|袋|盒|瓶|罐|支|条|块|卷|抽|听|颗|串|份|对|套|件|枝|朵|束|盆|根|桶|箱|杯|碗|张|'s)"
SPEC_RE = re.compile(r"[（(]?(?:约\s*)?(?:\d+\s*[-~－—]\s*)?\d[\d.]*\s*" + UNITS, re.IGNORECASE)
JISUDA_RE = re.compile(r"^(?:极.达|.速达)")
STOCK_WORDS = ("补货", "售罄", "预售", "已抢光", "到货")
FIXUPS = (("仅刺", "仅剩"), ("仅剰", "仅剩"), ("仅乘", "仅剩"), ("冰鮮", "冰鲜"))


def log(msg):
    print(f"{datetime.now():%H:%M:%S} {msg}", flush=True)


def now():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------- 截图与分析

def act(*args):
    r = subprocess.run([str(BIN / "act"), *map(str, args)], capture_output=True, text=True, timeout=30)
    return r.stdout.strip()


def get_window():
    """镜像窗口 (id, x, y, w, h, 是否在最前)；找不到返回 None。"""
    p = act("window").split()
    if len(p) != 6:
        return None
    return dict(id=int(p[0]), x=int(p[1]), y=int(p[2]), w=int(p[3]), h=int(p[4]), front=p[5] == "1")


def cursor():
    """(x, y, 距离上一次键盘按键的秒数)"""
    p = act("cursor").split()
    return (int(p[0]), int(p[1]), float(p[2]) if len(p) > 2 else 1e9) if len(p) >= 2 else None


def clean_name(t):
    t = re.sub(r"^[^一-龥A-Za-z0-9]+", "", t)
    return re.sub(r"[^一-龥A-Za-z0-9/&·]+$", "", t)


def best_entry(side, name):
    """左侧栏里最像 name 的那一项 (名字, y)：有一字不差的就用它，否则取最像的（至少六成像）。
    不能取「第一个像的」：秋季童装 / 秋季男装 / 秋季女装 互相都有七成半像。"""
    exact = next((t for t in side if t[0] == name), None)
    if exact:
        return exact
    scored = [(difflib.SequenceMatcher(None, t[0], name).ratio(), t) for t in side]
    scored = [x for x in scored if x[0] >= 0.6]
    return max(scored, key=lambda x: x[0])[1] if scored else None


def same_name(a, b):
    return bool(a and b) and (a == b or difflib.SequenceMatcher(None, a, b).ratio() >= 0.6)


class Page:
    """一张截图的分析结果，坐标统一换算到 383×834 的基准布局。"""

    def __init__(self, raw, path=None):
        self.path = path
        sx, sy = raw["width"] / BASE_W, raw["height"] / BASE_H
        self.lines = [dict(x=l["x"] / sx, y=l["y"] / sy, w=l["w"] / sx, h=l["h"] / sy, conf=l["conf"],
                           text=l["text"].strip(), fg=l["fg"], bg=l["bg"]) for l in raw["lines"]]
        self.bar_y = raw["barY"] / sy if raw.get("barY") is not None else None
        self.hash = raw["listHash"]

    def tab_y(self):
        ys = [l["y"] for l in self.lines if l["y"] > 755 and any(w in l["text"] for w in TAB_WORDS)]
        return min(ys) if len(ys) >= 2 else None

    def header(self):
        """列表表头那一行（极速达 / 全国配 / 销量 / 价格）。子分类顶部有活动横幅时它会被顶到下面，所以 y 范围放宽；
        用「销量」「价格」定位，避免把商品上的「全国配」角标当成表头。"""
        cand = [l for l in self.lines if 195 < l["y"] < 520 and l["x"] > 240 and "销量" in l["text"]]
        if not cand:
            price = [l for l in self.lines if 195 < l["y"] < 520 and l["x"] > 280 and l["text"].startswith("价格")]
            cand = [l for l in self.lines if 195 < l["y"] < 520 and l["x"] >= 100 and "全国配" in l["text"]
                    and any(abs(l["y"] - q["y"]) < 10 for q in price)]
        return min(cand, key=lambda l: l["y"]) if cand else None

    def top_selected(self):
        """顶部大类条里选中的那个（蓝底白字）。"""
        for l in self.lines:
            r, g, b = l["bg"]
            if 160 < l["y"] < 200 and b > 140 and r < 90 and b - r > 80:
                return clean_name(l["text"])
        return None

    def top_labels(self):
        return [l for l in self.lines if 160 < l["y"] < 200 and l["x"] < 330 and len(clean_name(l["text"])) >= 2]

    def panel_open(self):
        """顶部「展开」后的「全部分类」面板是否开着。"""
        return (any(l["text"].startswith("全部分类") and l["y"] < 200 for l in self.lines)
                and any("收起" in l["text"] for l in self.lines))

    def panel_labels(self):
        bottom = min((l["y"] for l in self.lines if "收起" in l["text"]), default=560)
        return [l for l in self.lines if 150 < l["y"] < min(bottom, 560) - 5 and len(clean_name(l["text"])) >= 2
                and "全部分类" not in l["text"]]

    def panel_selected(self):
        for l in self.panel_labels():
            r, g, b = l["bg"]
            if b > 140 and r < 90 and b - r > 80:
                return clean_name(l["text"])
        return None

    def sidebar(self):
        out = []
        for l in self.lines:
            if l["x"] + l["w"] < SIDEBAR_X_MAX + 8 and 205 < l["y"] < 745:
                name = clean_name(l["text"])
                if len(name) >= 2:
                    out.append((name, l["y"] + l["h"] / 2))
        return sorted(out, key=lambda t: t[1])

    def sub_selected(self):
        if self.bar_y is None:
            return None
        best = min(self.sidebar(), key=lambda t: abs(t[1] - self.bar_y), default=None)
        return best[0] if best and abs(best[1] - self.bar_y) <= 16 else None

    def base_ok(self):
        """是不是山姆「分类」页（底部标签栏 + 顶部选中的大类都在）。"""
        return self.tab_y() is not None and self.top_selected() is not None

    def ok(self):
        """分类页，并且右边是带表头的商品列表。"""
        return self.base_ok() and self.header() is not None

    def net_error(self):
        """列表区显示「网络异常，请尝试刷新」时，返回「重试」按钮那行；否则 None。"""
        area = [l for l in self.lines if l["x"] >= IMG_X_MIN and 300 < l["y"] < 745]
        if not any("网络异常" in l["text"] for l in area):
            return None
        return next((l for l in area if l["text"].strip() == "重试"), None)

    def oos_footer(self):
        """子分类末尾那行「点击展开更多无货商品」。"""
        for l in self.lines:
            if l["x"] >= IMG_X_MIN and 230 < l["y"] < 745 and "无货商品" in l["text"] and "展开" in l["text"]:
                return l
        return None


class Ocr:
    """常驻的识别进程（bin/analyze --serve）。每页新开进程的话，系统有时每次都要重新准备识别模型，
    一页要等二三十秒；常驻后只在第一页慢一次。"""

    proc = None

    @classmethod
    def ask(cls, path, timeout=180):
        import select
        if cls.proc is None or cls.proc.poll() is not None:
            cls.proc = subprocess.Popen([str(ANALYZE), "--serve", "3"], stdin=subprocess.PIPE,
                                        stdout=subprocess.PIPE, stderr=subprocess.DEVNULL)
        try:
            cls.proc.stdin.write((str(path) + "\n").encode("utf-8"))
            cls.proc.stdin.flush()
            buf, deadline = b"", time.time() + timeout
            while not buf.endswith(b"\n"):
                left = deadline - time.time()
                if left <= 0 or not select.select([cls.proc.stdout], [], [], left)[0]:
                    raise TimeoutError
                chunk = os.read(cls.proc.stdout.fileno(), 1 << 16)
                if not chunk:
                    raise BrokenPipeError
                buf += chunk
            return json.loads(buf)
        except (TimeoutError, BrokenPipeError, OSError, ValueError):
            cls.close()                         # 出问题就丢掉这个进程，下一页重新起
            return None

    @classmethod
    def close(cls):
        if cls.proc is not None:
            cls.proc.kill()
            cls.proc = None


OCR_CACHE = Path.home() / "Library" / "Caches" / "analyze"


def analyze_raw(path):
    """识别一张截图，返回 bin/analyze 的原始结果（dict）；识别失败返回 None。"""
    t = time.time()
    raw = Ocr.ask(path)
    if raw is not None and "error" in raw:
        # 系统给识别程序建的模型缓存坏了（在命令沙箱里跑过一次就会留下空的缓存条目，之后每张图都失败）。
        # 清掉缓存、重启识别进程再试一次；重建缓存的那一张会慢二三十秒。
        import shutil
        Ocr.close()
        shutil.rmtree(OCR_CACHE, ignore_errors=True)
        log("  ⚠ 识别模型缓存损坏，已清理并重建")
        raw = Ocr.ask(path)
    if time.time() - t > 10:
        log(f"  ⚠ 这一页识别用了 {time.time() - t:.0f} 秒（识别模型刚加载时会这样）")
    if raw is None or "error" in raw:
        return None
    return raw


def analyze(path):
    raw = analyze_raw(path)
    return None if raw is None else Page(raw, path)


def snap(win, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    r = subprocess.run(["screencapture", "-x", "-o", "-l", str(win["id"]), str(path)], capture_output=True, timeout=30)
    if r.returncode != 0 or not path.exists() or path.stat().st_size < 5000:
        return None
    return analyze(path)


# ---------------------------------------------------------------- 解析商品卡

def kind(l):
    fr, fg, fb = l["fg"]
    br, bg, bb = l["bg"]
    neutral_bg = abs(br - bg) < 14 and abs(bg - bb) < 14
    if min(br, bg, bb) < 228 and not neutral_bg:
        return "badge"                     # 彩色底：极速达、NEW、榜单等角标
    mx, mn = max(l["fg"]), min(l["fg"])
    if mx < 118:
        return "black"                     # 标题
    if mx - mn < 32:
        return "gray"                      # 卖点
    if fr > 195 and fg < 90 and fb < 95:
        return "red"                       # 价格
    return "color"                         # 冷冻/冰鲜、评价等标签


def price_of(l):
    t = re.sub(r"(?<=\d)[,，](?=\d{3})", "", l["text"].replace(" ", ""))
    m = PRICE_RE.match(t)
    return float(m.group(1)) if m else None


def fix(t):
    for a, b in FIXUPS:
        t = t.replace(a, b)
    return t


FLOAT_ZONE = (652, 722, 316)         # 右下角悬浮按钮占的位置：y 上界、y 下界、x 左界


def parse_cards(page, at_top=False, strict=True):
    """返回本页「完整可见」的商品卡。被顶部、底部截断或被悬浮按钮挡住标题的卡不要，
    它们会在相邻页的别的位置完整出现。strict=False 时不排除悬浮按钮区域（列表到底、
    这些卡再也挪不开时用）。"""
    hdr, tab = page.header(), page.tab_y()
    if hdr is None or tab is None:
        return []
    top, bottom = hdr["y"] + hdr["h"], tab - 28
    inside = [l for l in page.lines if l["y"] >= top - 2 and l["y"] + l["h"] <= bottom + 3]
    L = sorted([l for l in inside if LIST_X_MIN <= l["x"] < LIST_X_MAX], key=lambda l: (round(l["y"] / 4), l["x"]))
    stock = [l for l in inside if IMG_X_MIN <= l["x"] < LIST_X_MIN and any(w in l["text"] for w in STOCK_WORDS)]

    def is_price(l):
        return kind(l) == "red" and l["x"] < 206 and l["h"] >= 12 and price_of(l) is not None

    groups, cur = [], []
    for l in L:
        cur.append(l)
        if is_price(l):
            groups.append(cur)
            cur = []
    cards = []
    for i, g in enumerate(groups):
        t0 = next((l for l in g if kind(l) == "black"), None)
        if t0 is None:
            continue                                   # 标题已滚出顶部
        pre = [l for l in g if l["y"] + l["h"] / 2 < t0["y"] - 3]
        if pre and cards and i > 0:                    # 价格右边的单位价、价格下面的榜单条，归上一张卡
            cards[-1]["tags"] += ["单位价 " + l["text"] if kind(l) == "gray" else fix(l["text"]) for l in pre
                                  if kind(l) != "gray" or re.match(r"^[¥￥]\s*\d", l["text"])]
        if i == 0 and t0["y"] - top < (8 if at_top else 24):
            continue                                   # 第一行标题可能被表头盖住
        body = [l for l in g if l not in pre]
        line1 = sorted([l for l in body if kind(l) == "black" and abs(l["y"] - t0["y"]) <= 6], key=lambda l: l["x"])
        line2 = sorted([l for l in body if kind(l) == "black" and 10 < l["y"] - t0["y"] < 27], key=lambda l: l["x"])
        tl = line1 + line2
        if strict and any(l["y"] + l["h"] > FLOAT_ZONE[0] and l["y"] < FLOAT_ZONE[1] and l["x"] + l["w"] > FLOAT_ZONE[2]
                          for l in tl):
            continue                                   # 标题伸进了悬浮按钮的位置，可能被挡住
        title = " ".join(l["text"] for l in line1) + "".join(l["text"] for l in line2)
        tags = []
        jisuda = any(kind(l) == "badge" and "达" in l["text"] for l in body)
        if line1[0]["x"] < 206 and JISUDA_RE.match(title):
            title, jisuda = title[3:], True            # 角标被 OCR 并进了标题
        title = re.sub(r"^[^一-龥A-Za-z0-9（(]+", "", title).strip()
        if title.startswith("仅门店可售"):
            title, tags = title[5:].strip(), ["仅门店可售"]
        if len(title) < 2:
            continue
        price_line = g[-1]
        used = set(map(id, tl + [price_line]))
        unit = [l for l in body if kind(l) == "gray" and id(l) not in used and re.match(r"^[¥￥]\s*\d", l["text"])]
        used |= set(map(id, unit))              # 价格右边的灰色小字是单位价，如「¥7.49/份」
        desc = "".join(l["text"] for l in body if kind(l) == "gray" and id(l) not in used)
        tags += ["单位价 " + l["text"] for l in unit]
        tags += [fix(l["text"]) for l in body if kind(l) in ("color", "badge", "red") and id(l) not in used
                 and not ("达" in l["text"] and len(l["text"]) <= 4)]
        y0, y1 = t0["y"] - 14, price_line["y"] + price_line["h"] + 6
        tags += [l["text"] for l in stock if y0 <= l["y"] <= y1]
        cards.append(dict(title=title, price=price_of(price_line), desc=desc, tags=tags, jisuda=jisuda,
                          conf=round(min(l["conf"] for l in tl + [price_line]), 2), y=round(t0["y"])))
    return cards


def split_spec(title):
    """「商品名 + 规格」拆开。拆不出来就规格留空，原标题另有一列保留。"""
    m = SPEC_RE.search(title)
    if not m or len(re.sub(r"[^一-龥A-Za-z]", "", title[:m.start()])) < 2:
        return title, ""
    return title[:m.start()].strip(" ，,"), title[m.start():].strip()


def norm(t):
    return re.sub(r"[\s·,，.。:：;；|｜\-—…*×xX()（）'‘’\"“”]", "", t).lower()


# ---------------------------------------------------------------- 状态与结果

def load_state():
    if STATE.exists():
        return json.loads(STATE.read_text(encoding="utf-8"))
    return {"tops": {}}


def save_state(state):
    DATA.mkdir(parents=True, exist_ok=True)
    tmp = STATE.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, STATE)


def read_records():
    if not ITEMS.exists():
        return []
    return [json.loads(x) for x in ITEMS.read_text(encoding="utf-8").splitlines() if x.strip()]


def merged_items():
    """items.jsonl 是只追加的流水：商品记录 + 之后的「改子分类」记录。这里合成最终的商品列表。"""
    out, idx = [], {}
    for r in read_records():
        if r.get("_relabel"):
            it = idx.get(r["id"])
            if it and r["sub"] != it["sub"]:
                if it["sub"] and it["sub"] not in it["also"]:
                    it["also"].append(it["sub"])
                it["sub"] = r["sub"]
                if r["sub"] in it["also"]:
                    it["also"].remove(r["sub"])
            continue
        r.setdefault("also", [])
        idx[r["id"]] = r
        out.append(r)
    return out


class Store:
    """已抓到的商品。同一商品会在相邻两页、不同子分类里重复出现，这里去重。"""

    def __init__(self):
        self.items = merged_items()
        self.next_id = max((it["id"] for it in self.items), default=0) + 1
        self.index = {}
        for it in self.items:
            self.index.setdefault((it["top"], it["price"]), []).append(it)
        self.fh = None

    def find(self, top, card):
        n = norm(card["title"])
        for it in self.index.get((top, card["price"]), []):
            m = norm(it["title"])
            if m == n or difflib.SequenceMatcher(None, m, n).ratio() >= 0.8:
                return it
        return None

    def add(self, top, sub, card, shot, relabel=True):
        """新商品返回 True。已有的商品：relabel 时把子分类改成现在这个（原来的记到「其他子分类」）。"""
        it = self.find(top, card)
        if it:
            if relabel and sub and sub != it["sub"]:
                self._write(dict(_relabel=True, id=it["id"], sub=sub))
                if it["sub"] and it["sub"] not in it["also"]:
                    it["also"].append(it["sub"])
                it["sub"] = sub
            return False
        rec = dict(id=self.next_id, top=top, sub=sub or "", title=card["title"], price=card["price"],
                   desc=card["desc"], tags=card["tags"], jisuda=card["jisuda"], conf=card["conf"], shot=shot,
                   at=now(), also=[], oos=bool(card.get("oos")))
        self.next_id += 1
        self._write(rec)
        self.items.append(rec)
        self.index.setdefault((top, rec["price"]), []).append(rec)
        return True

    def _write(self, rec):
        if self.fh is None:
            DATA.mkdir(parents=True, exist_ok=True)
            self.fh = open(ITEMS, "a", encoding="utf-8")
        self.fh.write(json.dumps(rec, ensure_ascii=False) + "\n")
        self.fh.flush()
        os.fsync(self.fh.fileno())


def export():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    items = merged_items()
    state = load_state()
    order = {}
    for top, st in state.get("tops", {}).items():
        for i, s in enumerate(st.get("subs", [])):
            order[(top, s)] = i
    tops = list(state.get("tops", {}))
    items.sort(key=lambda it: (tops.index(it["top"]) if it["top"] in tops else 99, order.get((it["top"], it["sub"]), 999), it["id"]))
    # App 是按大类抓的，同一个商品会挂在几个大类下（个护美妆 + 全球购、乳品烘焙 + 营养早餐……），各抓到一次。
    # 标题和价格都一样的只留一行（排在前面的那个大类），别的分类记在「其他所在分类」里；极速达 / 无货有一处是就算是。
    uniq, seen = [], {}
    for it in items:
        key = (norm(it["title"]), it["price"])
        if key in seen:
            first = seen[key]
            first["dup"].append(f"{it['top']}/{it['sub']}")
            first["jisuda"] = first.get("jisuda") or it.get("jisuda")
            first["oos"] = first.get("oos") or it.get("oos")
            first["tags"] = first.get("tags", []) + [t for t in it.get("tags", []) if t not in first.get("tags", [])]
            continue
        seen[key] = dict(it, dup=[], tags=list(it.get("tags", [])))
        uniq.append(seen[key])

    wb = Workbook()
    ws = wb.active
    ws.title = "商品"
    ws.append(["大类", "子分类", "商品名", "规格", "价格", "极速达", "无货", "标签", "卖点", "其他出现的子分类", "其他所在分类",
               "商品标题原文", "识别置信度", "截图", "抓取时间"])
    for it in uniq:
        name, spec = split_spec(it["title"])
        out_of_stock = it.get("oos") or any(w in t for t in it.get("tags", []) for w in STOCK_WORDS)
        ws.append([it["top"], it["sub"], name, spec, it["price"], "是" if it.get("jisuda") else "",
                   "是" if out_of_stock else "", " | ".join(it.get("tags", [])), it.get("desc", ""), "、".join(it.get("also", [])),
                   "、".join(it["dup"]), it["title"], it.get("conf", ""), it.get("shot", ""), it.get("at", "")])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, wd in enumerate([10, 12, 34, 18, 9, 7, 6, 30, 36, 18, 24, 40, 10, 22, 19], 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        row[0].number_format = "0.00"

    ws2 = wb.create_sheet("进度")
    ws2.append(["大类", "状态", "已翻页数", "抓到的条数（含跨大类重复）", "当前子分类", "已抓完的子分类", "更新时间"])
    for top, st in state.get("tops", {}).items():
        n = sum(1 for it in items if it["top"] == top)
        ws2.append([top, {"done": "已完成", "running": "未完成", "skipped": "已跳过"}.get(st.get("status"), st.get("status", "")),
                    st.get("pages", 0), n, st.get("current_sub", ""), "、".join(st.get("done_subs", [])), st.get("updated", "")])
    for c in ws2[1]:
        c.font = Font(bold=True)
    for i, wd in enumerate([12, 10, 10, 26, 14, 90, 20], 1):
        ws2.column_dimensions[get_column_letter(i)].width = wd
    # 京东「山姆会员商店官方旗舰店」的商品（jd.py / jd_receiver.py 抓的），单独一个 sheet。价格只有山姆会员价。
    try:
        import jd
        jd_rows = jd.jd_items()
    except Exception:                       # noqa: BLE001  没抓过京东就没有这个 sheet
        jd_rows = []
    if jd_rows:
        ws3 = wb.create_sheet("京东山姆店")
        ws3.append(["一级分类", "二级分类", "商品名", "山姆会员价", "京东商品编号", "商品链接", "图片文件", "图片地址",
                    "评价数", "其他所在分类", "抓取时间"])
        for it in jd_rows:
            try:
                price = float(it["sam"]) if it.get("sam") else None
            except ValueError:
                price = it["sam"]
            pic = jd.image_path(it)
            ws3.append([it.get("top", ""), it.get("sub", ""), it["name"], price, it["sku"],
                        f"https://item.jd.com/{it['sku']}.html",
                        f"jd_images/{pic.name}" if pic.exists() else "",
                        jd.IMG_HOST + it["img"] if str(it.get("img", "")).startswith("jfs/") else "",
                        it.get("ev", "").replace("条评价", ""),
                        "、".join(f"{a}/{b}" for a, b in it.get("also", [])), it.get("at", "")])
        for c in ws3[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
            c.alignment = Alignment(horizontal="center")
        ws3.freeze_panes = "A2"
        ws3.auto_filter.ref = ws3.dimensions
        for i, wd in enumerate([12, 16, 48, 12, 18, 38, 26, 40, 10, 24, 19], 1):
            ws3.column_dimensions[get_column_letter(i)].width = wd
        for row in ws3.iter_rows(min_row=2, min_col=4, max_col=5):
            row[0].number_format = "0.00"
            row[1].number_format = "@"          # 商品编号按文本，避免变成科学计数法

    # App 和京东两边对比（compare.py）：各自独有的商品各一个 sheet。两边没有共同的商品编号，是按名字、价格、规格认的，
    # 所以每行带上对面最像的那个商品和相似度——相似度高却没认上的，多半是同名不同包装 / 不同款，拿不准的可以人工看一眼。
    if jd_rows:
        import compare
        app_only, jd_only, _ = compare.compare(uniq, jd_rows)       # 用去重后的 App 商品比
        ws4 = wb.create_sheet("App有京东无")
        ws4.append(["大类", "子分类", "商品名", "规格", "App价格", "其他所在分类", "京东最像的商品", "京东山姆会员价", "相似度", "备注"])
        for it, near, s in app_only:
            name, spec = split_spec(it["title"])
            ws4.append([it["top"], it["sub"], name, spec, it["price"], "、".join(it["dup"]), near["name"] if near else "",
                        compare.jd_price(near) if near else None, round(s, 2) if near else None,
                        compare.note(s, it["price"], compare.jd_price(near)) if near else ""])
        ws5 = wb.create_sheet("京东有App无")
        ws5.append(["一级分类", "二级分类", "商品名", "山姆会员价", "京东商品编号", "商品链接", "App最像的商品", "App价格", "相似度", "备注"])
        for it, near, s in jd_only:
            ws5.append([it.get("top", ""), it.get("sub", ""), it["name"], compare.jd_price(it), it["sku"],
                        f"https://item.jd.com/{it['sku']}.html", near["title"] if near else "",
                        near["price"] if near else None, round(s, 2) if near else None,
                        compare.note(s, near["price"], compare.jd_price(it)) if near else ""])
        for w_, widths, money in ((ws4, [10, 12, 34, 18, 9, 22, 44, 14, 8, 34], (5, 8)), (ws5, [12, 16, 48, 12, 18, 38, 40, 9, 8, 34], (4, 8))):
            for c in w_[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = PatternFill("solid", fgColor="1F4E78")
                c.alignment = Alignment(horizontal="center")
            w_.freeze_panes = "A2"
            w_.auto_filter.ref = w_.dimensions
            for i, wd in enumerate(widths, 1):
                w_.column_dimensions[get_column_letter(i)].width = wd
            for col in money:
                for row in w_.iter_rows(min_row=2, min_col=col, max_col=col):
                    row[0].number_format = "0.00"
        for row in ws5.iter_rows(min_row=2, min_col=5, max_col=5):
            row[0].number_format = "@"              # 商品编号按文本

    tmp = XLSX.with_name("tmp_" + XLSX.name)
    wb.save(tmp)
    os.replace(tmp, XLSX)
    return len(uniq)


# ---------------------------------------------------------------- 界面操作

class Stop(Exception):
    """整个抓取停下来（进度已存）。"""


class SkipTop(Exception):
    """这个大类进不去或结构不同，跳过它继续下一个。"""


class NoList(Exception):
    """在分类页上，但当前子分类右边不是普通的商品列表。"""


# App 里的 18 个大类。不想抓的大类写进 data/exclude.txt（一行一个）；每开始一个大类前都会重新读这个文件，
# 所以抓取过程中往里加名字也生效，不用重启。
ALL_TOPS = ["肉蛋果蔬", "乳品烘焙", "速食冷冻", "休闲零食", "酒水饮料", "粮油干货", "个护美妆", "母婴玩具", "全球购",
            "家清纸品", "家电家居", "服饰家纺", "营养保健", "萌宠生活", "眼镜助听", "线上专享", "礼品卡", "营养早餐"]
EXCLUDE = DATA / "exclude.txt"


def excluded():
    if not EXCLUDE.exists():
        return set()
    return {x.strip() for x in EXCLUDE.read_text(encoding="utf-8").splitlines() if x.strip()}


class Runner:
    def __init__(self, top, max_pages, wait_minutes):
        self.top, self.max_pages, self.wait_minutes = top, max_pages, wait_minutes
        self.state = load_state()
        self.store = Store()
        self.win = get_window()
        self.last_cursor = None
        self.last_act = None                    # 脚本上一次动鼠标的时刻
        self.resync = False
        self.only_sub = None
        self.session_pages = 0
        if self.win is None:
            raise Stop("没找到 iPhone 镜像窗口，请先打开并连接手机")

    # 基准坐标 → 屏幕坐标
    def sx(self, x):
        return round(self.win["x"] + x * self.win["w"] / BASE_W)

    def sy(self, y):
        return round(self.win["y"] + y * self.win["h"] / BASE_H)

    def guard(self):
        """动鼠标之前：有没有 STOP 文件、光标有没有被人挪走、窗口还在不在。"""
        if STOP.exists():
            STOP.unlink()
            raise Stop("收到 STOP 文件")
        c = cursor()
        if self.last_cursor and c and (abs(c[0] - self.last_cursor[0]) > 40 or abs(c[1] - self.last_cursor[1]) > 40):
            self.wait_ready("检测到鼠标被移动（有人在用这台 Mac，或在替脚本点手机）")
        elif self.last_act and c and c[2] < time.time() - self.last_act - 0.3:
            self.wait_ready("检测到有人在用键盘")
        w = get_window()
        if w is None:
            return False
        self.win = w
        return True

    def _acted(self, out):
        m = re.search(r"done (-?\d+) (-?\d+)", out)
        self.last_cursor = (int(m.group(1)), int(m.group(2))) if m else None
        self.last_act = time.time()

    def scroll_list(self, dy):
        self._acted(act("scroll", self.sx(220), self.sy(500), dy, 14))

    def scroll_sidebar(self, dy):
        self._acted(act("scroll", self.sx(48), self.sy(480), dy, 12))

    def click(self, x, y, zone):
        """只允许点：顶部大类条、「展开」按钮、「全部分类」面板里的大类、左侧子分类栏、
        列表末尾的「点击展开更多无货商品」那行字、网络异常页上的「重试」。"""
        ok = ((zone == "sidebar" and x < SIDEBAR_X_MAX and 205 < y < 745)
              or (zone == "top" and x < 330 and 120 < y < 200)
              or (zone == "expand" and x > 335 and 130 < y < 200)
              or (zone == "panel" and 150 < y < 560)
              or (zone == "oos" and IMG_X_MIN < x < 330 and 230 < y < 745)
              or (zone == "retry" and IMG_X_MIN < x < 330 and 400 < y < 745))
        if not ok:
            raise Stop(f"拒绝点击范围外的位置 ({x:.0f},{y:.0f})")
        self._acted(act("click", self.sx(x), self.sy(y)))

    def tmp_shot(self):
        return snap(self.win, DATA / "shots" / "_current.png")

    def wait_ready(self, reason):
        """页面不对（来电话、弹窗、镜像断开）：只截图不动鼠标，等它恢复。"""
        log(f"⏸ 暂停：{reason}。现在可以放心用手机和这台 Mac；手机回到山姆「分类」页、"
            f"鼠标键盘都停 20 秒左右后自动继续（最多等 {self.wait_minutes} 分钟）")
        deadline, streak, last_c = time.time() + self.wait_minutes * 60, 0, None
        self.last_cursor, self.last_act, self.resync = None, None, True
        save_state(self.state)
        while time.time() < deadline:
            time.sleep(5)
            if STOP.exists():
                STOP.unlink()
                raise Stop("收到 STOP 文件")
            w = get_window()
            if w is None:
                streak = 0
                continue
            self.win = w
            p = self.tmp_shot()
            c = cursor()
            still = last_c is not None and c is not None and abs(c[0] - last_c[0]) <= 3 and abs(c[1] - last_c[1]) <= 3
            last_c = c
            ready = p is not None and (p.base_ok() or p.panel_open())
            idle = still and c[2] >= 20         # 鼠标没动、键盘 20 秒没按：人不在用这台 Mac 了
            streak = streak + 1 if (ready and idle) else 0
            if streak >= 3:
                log("▶ 页面已恢复，继续")
                return p
        raise Stop(f"等了 {self.wait_minutes} 分钟页面仍未恢复")

    def ready_page(self, need_header=True, nolist_ok=False):
        """拿到一张确认在「分类」页上的截图；不在就等。need_header 时还要求右边是带表头的商品列表，
        在分类页上却没有表头则抛 NoList（这个子分类不是普通列表）。"""
        def good(p):
            return p is not None and (p.ok() if need_header else p.base_ok())

        if not self.guard():
            return self.wait_ready("镜像窗口不见了")
        p = self.tmp_shot()
        for pause in (2.0, 2.5):                # 可能只是加载或提示条，多看两眼
            if good(p):
                break
            time.sleep(pause)
            p = self.tmp_shot()
        if good(p):
            return p
        if need_header and nolist_ok and p is not None and p.base_ok():
            raise NoList()
        return self.wait_ready("当前不是山姆「分类」页（可能来电话、弹窗或镜像断开）")

    def goto_sub(self, name):
        """在左侧栏点开某个子分类（列表会跳到它的开头）。看不到就上下滚动左侧栏找。"""
        for dy in (0, -260, -260, -260, -260, -260, -260, 300, 300, 300, 300, 300, 300, 300, 300):
            if dy:
                self.guard()
                self.scroll_sidebar(dy)
                time.sleep(1.0)
            p = self.ready_page(need_header=False)
            hit = best_entry(p.sidebar(), name)
            if not hit:
                continue
            for attempt in range(4):            # 点击偶尔会被吞掉：点完看高亮有没有落在这一项上，没有就再点
                self.guard()
                self.click(48, hit[1], "sidebar")
                time.sleep(2.4)
                p = self.ready_page(need_header=False)
                now_hit = best_entry(p.sidebar(), name)
                if now_hit and p.bar_y is not None and abs(now_hit[1] - p.bar_y) <= 16:
                    return True
                hit = now_hit or hit
            raise Stop(f"点了 4 次「{name}」高亮都没过去")
        return False

    def reveal_sub(self):
        """列表翻到了左侧栏当前看不见的子分类（左侧栏不会自己跟着滚）：把左侧栏滚到能看见高亮为止。"""
        for dy in (-220, -220, -220, -220, -220, -220, 260, 260, 260, 260, 260, 260, 260, 260, 260, 260):
            self.guard()
            self.scroll_sidebar(dy)
            time.sleep(1.0)
            p = self.ready_page(need_header=False)
            if self.resync:
                return None
            name = p.sub_selected()
            if name:
                return name
        return None

    def next_in_sidebar(self, target, done):
        """列表翻不动时，看左侧栏里当前子分类后面是否还有没抓的子分类；没有返回 None。"""
        for dy in (0, -200, -200):
            if dy:
                self.guard()
                self.scroll_sidebar(dy)
                time.sleep(1.0)
            p = self.ready_page(need_header=False)
            if self.resync:
                return None
            side = p.sidebar()
            names = [n for n, _ in side]
            me = best_entry(side, target)
            i = side.index(me) if me else None
            if i is not None and i + 1 < len(names):
                nxt = names[i + 1]
                return None if (nxt in done or nxt in SKIP_SUBS) else nxt
        return None

    def first_sub(self):
        """把左侧栏滚回顶部，返回第一个要抓的子分类名。"""
        for i in range(8):
            p = self.ready_page(need_header=False)
            names = [n for n, _ in p.sidebar()]
            if any(n in SKIP_SUBS for n in names[:3]) or i == 7:
                todo = [n for n in names if n not in SKIP_SUBS]
                return todo[0] if todo else None
            self.guard()
            self.scroll_sidebar(400)
            time.sleep(1.0)
        return None

    def select_top(self, top):
        """切到指定大类（None = 用当前选中的）。大类可能在顶部条里，也可能要到「全部分类」面板里点。"""
        for _ in range(5):
            if not self.guard():
                self.wait_ready("镜像窗口不见了")
                self.resync = False
                continue
            p = self.tmp_shot()
            if p is not None and p.panel_open():
                name = top or p.panel_selected()
                lab = next((l for l in p.panel_labels() if same_name(clean_name(l["text"]), name)), None)
                if lab is None:
                    raise SkipTop(f"「全部分类」里没找到大类「{name}」")
                self.click(lab["x"] + lab["w"] / 2, lab["y"] + lab["h"] / 2, "panel")
                time.sleep(3)
                continue
            p = self.ready_page(need_header=False)
            self.resync = False
            cur = p.top_selected()
            if top is None or same_name(cur, top):
                return cur
            lab = next((l for l in p.top_labels() if same_name(clean_name(l["text"]), top)), None)
            self.guard()
            if lab is not None:
                self.click(lab["x"] + lab["w"] / 2, lab["y"] + lab["h"] / 2, "top")
                time.sleep(3)
            else:
                self.click(360, 158, "expand")      # 顶部条里没有，打开「全部分类」面板
                time.sleep(1.8)
        raise SkipTop(f"切不到大类「{top}」")

    def run_one_sub(self):
        """只把一个子分类重扫一遍（补漏用），不改这个大类的完成状态。"""
        top = self.top or self.select_top(None)
        self.select_top(top)
        st = self.state["tops"].setdefault(top, dict(pages=0, subs=[], done_subs=[], current_sub="", started=now()))
        for k, v in (("pages", 0), ("subs", []), ("done_subs", []), ("skipped_subs", []), ("current_sub", "")):
            st.setdefault(k, v)
        keep = st["current_sub"]
        before = len(self.store.items)
        log(f"▶ 补扫「{top} / {self.only_sub}」（已有 {before} 个商品）")
        try:
            while self.scan_sub(top, st, self.only_sub) == "__resync__":
                pass
        except NoList:
            log(f"  ⚠ 「{self.only_sub}」右边不是普通商品列表")
        n = sum(1 for it in self.store.items if it["top"] == top and it["sub"] == self.only_sub)
        log(f"  ✓ 「{self.only_sub}」补扫完成：现有 {n} 个商品，本次新增 {len(self.store.items) - before} 个")
        st["current_sub"], st["updated"] = keep, now()
        save_state(self.state)

    # ------------------------------------------------------------ 主流程
    def run(self):
        if self.only_sub:
            return self.run_one_sub()
        if self.top and self.state["tops"].get(self.top, {}).get("status") == "done":
            return                              # 已抓完的大类不用再点进去
        cur = self.select_top(self.top)
        top = self.top or cur
        st = self.state["tops"].setdefault(top, dict(pages=0, subs=[], done_subs=[], current_sub="", started=now()))
        for k, v in (("pages", 0), ("subs", []), ("done_subs", []), ("skipped_subs", []), ("current_sub", "")):
            st.setdefault(k, v)
        if st.get("status") == "done":
            log(f"「{top}」已经抓完（{st['pages']} 页）。要重抓请删掉 data/state.json 里这一项")
            return
        st["status"] = "running"
        target = st.get("current_sub") or self.first_sub()
        if not target:
            raise SkipTop("左侧栏没识别出子分类")
        log(f"▶ 大类「{top}」，从子分类「{target}」{'继续' if st.get('current_sub') else '开始'}"
            f"（已有 {len(self.store.items)} 个商品、{st['pages']} 页）")
        while True:
            try:
                nxt = self.scan_sub(top, st, target)
            except NoList:
                log(f"  ⚠ 「{target}」右边不是普通商品列表，跳过这个子分类")
                if target not in st["skipped_subs"]:
                    st["skipped_subs"].append(target)
                nxt = self.next_in_sidebar(target, st["done_subs"])
                if self.resync:
                    self.resync = False
                    continue
            else:
                if nxt == "__resync__":
                    continue                    # 被打断后恢复：当前子分类从头再扫，去重保证不重复
                if nxt == "__limit__":
                    log(f"⏹ 本次已翻 {self.session_pages} 页（--max-pages），停在「{target}」")
                    break
                n = sum(1 for it in self.store.items if it["top"] == top and it["sub"] == target)
                m = sum(1 for it in self.store.items if it["top"] == top and it["sub"] == target and it.get("oos"))
                log(f"  ✓ 「{target}」完成，{n} 个商品（其中无货 {m}）· 累计 {len(self.store.items)} 个 · {st['pages']} 页")
            if target not in st["done_subs"]:
                st["done_subs"].append(target)
            if nxt is None:
                st["status"], st["current_sub"] = "done", ""
                log(f"✔ 列表到底了，大类「{top}」抓完")
                break
            if nxt in st["done_subs"] or nxt in SKIP_SUBS:
                raise Stop(f"「{target}」之后高亮跳到了已处理过的「{nxt}」，停下来请人看一眼")
            target = nxt
            st["current_sub"] = target
            save_state(self.state)
        st["updated"] = now()
        save_state(self.state)

    def scan_sub(self, top, st, target):
        """把一个子分类从头翻到尾。返回下一个子分类名；None=整个列表到底；
        "__resync__"=中途被打断过；"__limit__"=达到页数上限。"""
        if not self.goto_sub(target):
            raise Stop(f"左侧栏没找到子分类「{target}」")
        if self.resync:
            self.resync = False
            return "__resync__"
        st["current_sub"] = target
        if target not in st["subs"]:
            st["subs"].append(target)
        save_state(self.state)
        shots = DATA / "shots" / top
        shots.mkdir(parents=True, exist_ok=True)
        pending, prev_hash, prev_sig, still, first, lost = None, None, None, 0, True, False
        oos, oos_clicks = False, 0              # 是否已点开「无货商品」：点开之后新出现的商品都是无货的
        seen, barren = set(), 0                 # 连续很多页都没有新面孔 → 当作到底，防止原地空转

        def moved_on(p, cur):
            """高亮是不是已经离开当前子分类。名字相近的子分类（秋季童装/秋季男装）不能只靠名字像不像来判断：
            左侧栏里如果另有一项比高亮项更像当前子分类，说明高亮已经移走了。"""
            if cur == target:
                return False
            side = p.sidebar()
            if side and p.bar_y is not None:
                def score(n):
                    return 1.0 if n == target else difflib.SequenceMatcher(None, n, target).ratio()
                best = max(side, key=lambda t: score(t[0]))
                hl = min(side, key=lambda t: abs(t[1] - p.bar_y))
                if score(best[0]) >= 0.6:
                    return abs(best[1] - hl[1]) > 8
            return not same_name(cur, target)

        def same_view(a, b):
            """两页解析出的商品卡完全一样、位置也没动。到底后列表会回弹，像素指纹每次都略有不同，要靠这个判断。"""
            return bool(a) and b is not None and len(a) == len(b) and all(
                x["title"] == y["title"] and x["price"] == y["price"] and abs(x["y"] - y["y"]) <= 6 for x, y in zip(a, b))

        def flush(relabel, at_bottom=False):
            nonlocal pending
            if pending:
                page, was_first, shot, was_oos = pending
                cards = parse_cards(page, at_top=was_first, strict=not at_bottom)
                new = 0
                for c in cards:
                    c["oos"] = was_oos
                    new += self.store.add(top, target, c, shot, relabel)
                if st["pages"] % 10 == 0 or at_bottom:
                    log(f"    第 {st['pages']} 页 · {target} · 本页 {len(cards)} 个、新增 {new} 个 · 累计 {len(self.store.items)} 个")
            pending = None

        while True:
            p = self.ready_page(nolist_ok=first)
            if self.resync:
                self.resync = False
                flush(False)
                return "__resync__"
            btn = p.net_error()
            if btn is not None:
                # 手机上显示「网络异常，请尝试刷新」：点「重试」，然后把这个子分类从头再扫。
                self.net_tries = getattr(self, "net_tries", 0) + 1
                flush(False)
                if self.net_tries > 6:
                    self.net_tries = 0
                    self.wait_ready("手机上一直显示「网络异常」，点了 6 次重试都没好")
                else:
                    log(f"  ⚠ 「{target}」显示网络异常，点「重试」（第 {self.net_tries} 次）")
                    self.guard()
                    self.click(btn["x"] + btn["w"] / 2, btn["y"] + btn["h"] / 2, "retry")
                    time.sleep(3 + self.net_tries)
                self.resync = False
                return "__resync__"
            cur = p.sub_selected()
            if cur is None and p.bar_y is None and not lost:
                cur = self.reveal_sub()         # 高亮滚出了左侧栏的可见范围，把它找出来
                if self.resync:
                    self.resync = False
                    flush(False)
                    return "__resync__"
                lost = cur is None              # 实在找不到就按当前子分类继续，不反复找
            if cur and moved_on(p, cur):
                flush(True)
                self.net_tries = 0
                return cur                      # 高亮跳到下一个子分类：这一页不记，从它的开头重新扫
            sig = parse_cards(p, at_top=True, strict=False)     # 判断翻没翻动用宽松解析：屏幕上能看到的卡都算
            moved = p.hash != prev_hash and not same_view(sig, prev_sig)
            keys = {(c["title"], c["price"]) for c in sig}
            if moved:
                barren = barren + 1 if (keys <= seen) else 0
                seen |= keys
                if barren >= 30:
                    log(f"  ⚠ 「{target}」连续 30 页没有新商品，按到底处理")
                    moved, still = False, 3
            if os.environ.get("SAMS_DEBUG"):
                log(f"    · 页面 hash={p.hash[:8]} {'动了' if moved else '没动'} 高亮={cur} 卡片={len(sig)}")
            if not moved:
                still += 1
                if still >= 3:
                    # 翻不动了。每个子分类是独立的列表（末尾写着「上滑，查看下一分类」），所以这一屏都属于当前子分类。
                    # 可能是整个大类到底，也可能只是上滑没触发切换：看左侧栏后面还有没有子分类。
                    flush(True, at_bottom=True)
                    self.net_tries = 0
                    nxt = self.next_in_sidebar(target, st["done_subs"])
                    if self.resync:
                        self.resync = False
                        return "__resync__"
                    return nxt
                time.sleep(0.8 + still)         # 可能在加载下一段
            else:
                flush(True)
                still = 0
                st["pages"] += 1
                name = f"{st['pages']:05d}.png"
                os.replace(p.path, shots / name)
                pending = (p, first, f"{top}/{name}", oos)
                st["updated"] = now()
                save_state(self.state)
                prev_hash, prev_sig = p.hash, sig
                self.session_pages += 1
                if self.session_pages >= self.max_pages:
                    flush(True)
                    return "__limit__"
                foot = p.oos_footer()
                if foot is not None and oos_clicks < 3:
                    # 「点击展开更多无货商品」出现时，它上面的有货商品都已在屏幕上：先全部记下，再点开。
                    flush(True, at_bottom=True)
                    self.guard()
                    self.click(foot["x"] + foot["w"] / 2, foot["y"] + foot["h"] / 2, "oos")
                    oos, oos_clicks, first = True, oos_clicks + 1, False
                    time.sleep(1.8)
                    continue                    # 展开后原地再看，不滚动
            first = False
            self.guard()
            self.scroll_list(SCROLL_DY - random.randint(0, 30))
            time.sleep(1.0 + random.random() * 0.6)


# ---------------------------------------------------------------- 命令

def cmd_run(a):
    code, r = 0, None
    try:
        r = Runner(None, a.max_pages, a.wait_minutes)
        r.only_sub = a.only_sub
        tops = ALL_TOPS if a.top == "all" else (a.top.split(",") if a.top else [None])
        skips = 0
        for top in tops:                        # --top 可以是 all，或用逗号给多个大类，依次抓
            r.top = top.strip() if top else None
            if r.top in excluded():
                log(f"「{r.top}」在排除名单里（data/exclude.txt），不抓")
                continue
            try:
                r.run()
                skips = 0
            except SkipTop as e:
                log(f"⚠ 跳过大类「{r.top}」：{e}")
                st = r.state["tops"].setdefault(r.top, dict(pages=0, subs=[], done_subs=[], current_sub=""))
                st["status"], st["note"] = "skipped", str(e)
                save_state(r.state)
                skips += 1
                if skips >= 2:
                    raise Stop("连续两个大类都进不去，可能是手机状态不对")
            if r.session_pages >= r.max_pages:
                break
    except Stop as e:
        log(f"⏹ 已停止：{e}。进度已保存，重新执行 run 即从断点继续")
        code = 2
    except KeyboardInterrupt:
        log("⏹ 手动中断。进度已保存")
        code = 2
    finally:
        if r is not None:
            save_state(r.state)
    n = export()
    log(f"Excel 已更新：{XLSX}（{n} 个商品）")
    return code


def cmd_status(a):
    state, items = load_state(), merged_items()
    for top, st in state.get("tops", {}).items():
        n = sum(1 for it in items if it["top"] == top)
        print(f"{top}: {st.get('status')} · {st.get('pages', 0)} 页 · {n} 个商品 · 当前子分类「{st.get('current_sub')}」 · {st.get('updated')}")
        print("  已抓完：" + "、".join(st.get("done_subs", [])))
        if st.get("skipped_subs") or st.get("note"):
            print("  跳过：" + "、".join(st.get("skipped_subs", [])) + " " + st.get("note", ""))
    if not state.get("tops"):
        print("还没有抓取记录")
    return 0


def cmd_rescan(a):
    """把大类标回「未完成」，下次 run 时从头再扫一遍（已有商品保留并去重，只补新出现的）。"""
    state = load_state()
    for top, st in state.get("tops", {}).items():
        if a.top in (None, "all") or top in a.top.split(","):
            st.update(status="running", current_sub="", done_subs=[], skipped_subs=[])
            print("已标回未完成：", top)
    save_state(state)
    return 0


def cmd_export(a):
    print(f"{export()} 个商品 → {XLSX}")
    return 0


def cmd_replay(a):
    seen = set()
    for f in sorted(Path(a.dir).glob("*.png")):
        p = analyze(f)
        if p is None or not p.ok():
            print(f"== {f.name} 不是分类页，跳过")
            continue
        print(f"== {f.name} top={p.top_selected()} sub={p.sub_selected()} hash={p.hash[:8]}")
        for c in parse_cards(p):
            seen.add((norm(c["title"]), c["price"]))
            name, spec = split_spec(c["title"])
            print(f"   {c['price']:>8.2f}  {name} 【{spec}】 {'极速达 ' if c['jisuda'] else ''}{' | '.join(c['tags'])}  conf={c['conf']}")
    print("去重后商品数：", len(seen))
    return 0


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run")
    r.add_argument("--top", help="大类名；all = 全部 18 个大类；可用逗号给多个；默认用手机上当前选中的大类")
    r.add_argument("--only-sub", help="只把 --top 这个大类里的这一个子分类重扫一遍（补漏用）")
    r.add_argument("--max-pages", type=int, default=100000)
    r.add_argument("--wait-minutes", type=int, default=30)
    r.set_defaults(fn=cmd_run)
    sub.add_parser("status").set_defaults(fn=cmd_status)
    rs = sub.add_parser("rescan")
    rs.add_argument("--top")
    rs.set_defaults(fn=cmd_rescan)
    sub.add_parser("export").set_defaults(fn=cmd_export)
    rp = sub.add_parser("replay")
    rp.add_argument("dir")
    rp.set_defaults(fn=cmd_replay)
    a = ap.parse_args()
    sys.exit(a.fn(a))


if __name__ == "__main__":
    main()
