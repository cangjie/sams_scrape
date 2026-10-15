import sys
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import compare  # noqa: E402
import jd  # noqa: E402
import scrape  # noqa: E402


class RegressionTest(unittest.TestCase):
    """抽函数以后，商品去重和京东配对的结果要和 data/山姆商品_暂存.xlsx（2026-10-15 导出）一致。"""

    @classmethod
    def setUpClass(cls):
        from openpyxl import load_workbook
        wb = load_workbook(scrape.XLSX, read_only=True)
        cls.sheet_names = [r[2] for r in wb["商品"].iter_rows(min_row=2, values_only=True)]
        cls.app_only_names = [r[2] for r in wb["App有京东无"].iter_rows(min_row=2, values_only=True)]
        cls.products = scrape.unique_items()
        cls.jds = jd.jd_items()

    def test_unique_items_same_as_excel(self):
        names = [scrape.split_spec(p["title"])[0] for p in self.products]
        self.assertEqual(len(names), 4604)
        self.assertEqual(names, self.sheet_names)

    def test_members_cover_all_records(self):
        self.assertEqual(sum(len(p["members"]) for p in self.products), len(scrape.merged_items()))
        for p in self.products:
            self.assertEqual(p["members"][0]["id"], p["id"])

    def test_match_same_as_excel(self):
        pairs, _, _ = compare.match(self.products, self.jds)
        self.assertEqual(len(pairs), 745)
        hit = {i for i, _, _ in pairs}
        self.assertEqual(len(hit), 739)
        app_only = [scrape.split_spec(p["title"])[0] for i, p in enumerate(self.products) if i not in hit]
        self.assertEqual(app_only, self.app_only_names)
        self.assertTrue(all(0.15 <= s <= 1 for _, _, s in pairs))

    def test_compare_unchanged(self):
        app_only, jd_only, n = compare.compare(self.products, self.jds)
        self.assertEqual((len(app_only), len(jd_only), n), (3865, 17, 745))


if __name__ == "__main__":
    unittest.main()
